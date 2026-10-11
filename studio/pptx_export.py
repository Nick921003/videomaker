"""將 Deck v1 畫為原生可編輯 PowerPoint 投影片。"""

from io import BytesIO
import math
import unicodedata

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from studio.deck_contracts import DeckError, canonical_deck


_BG = RGBColor(250, 246, 237)
_INK = RGBColor(47, 52, 45)
_ACCENT = RGBColor(169, 72, 42)
_SOFT = RGBColor(240, 229, 211)
_WHITE = RGBColor(255, 255, 255)


def _text_units(value, mono=False):
	total = 0.0
	for character in value:
		if character == "\t":
			total += 2.4
		elif unicodedata.east_asian_width(character) in ("F", "W"):
			total += 1.0
		elif character.isspace():
			total += 0.32
		else:
			total += 0.62 if mono else 0.6
	return total


def _fitted_size(content, width, height, preferred, minimum, mono=False):
	"""保守估算行數；無法在最低字級容納時拒絕匯出。"""
	usable_width = max(1, (width - 0.08) * 72 * 0.88)
	usable_height = max(1, (height - 0.04) * 72)
	for size in range(preferred, minimum - 1, -1):
		lines = sum(max(1, math.ceil(_text_units(line, mono) * size / usable_width))
			for line in content.split("\n"))
		if lines * size * 1.22 <= usable_height:
			return size
	raise DeckError("簡報文字過長，無法以最低可讀字級排入投影片")


def _text(slide, content, x, y, width, height, size=20, min_size=None, color=_INK, bold=False, font="Aptos", shape=None):
	shape = shape or slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
	shape.left, shape.top = Inches(x), Inches(y)
	shape.width, shape.height = Inches(width), Inches(height)
	chosen_size = _fitted_size(content, width, height, size,
		min_size if min_size is not None else max(10, math.ceil(size * 0.75)),
		mono=font == "Consolas")
	frame = shape.text_frame
	frame.word_wrap = True
	frame.auto_size = MSO_AUTO_SIZE.NONE
	frame.margin_left = frame.margin_right = Inches(0.04)
	frame.margin_top = frame.margin_bottom = Inches(0.02)
	frame.text = content
	for paragraph in frame.paragraphs:
		for run in paragraph.runs:
			run.font.name = font
			run.font.size = Pt(chosen_size)
			run.font.bold = bold
			run.font.color.rgb = color
	return shape


def _panel(slide, x, y, width, height, fill=_SOFT):
	shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
		Inches(x), Inches(y), Inches(width), Inches(height))
	shape.fill.solid()
	shape.fill.fore_color.rgb = fill
	shape.line.fill.background()
	return shape


def _draw_vertical_list(slide, block, self_study):
	top, step = (1.95, 0.69) if self_study else (2.08, 0.83)
	for index, item in enumerate(block["items"]):
		y = top + index * step
		_panel(slide, 0.8, y, 11.7, 0.60 if self_study else 0.68)
		prefix = f"{index + 1}." if block["type"] == "steps" else "•"
		_text(slide, prefix, 1.05, y + (0.09 if self_study else 0.11),
			0.45, 0.40 if self_study else 0.45, size=18 if self_study else 21, color=_ACCENT, bold=True)
		_text(slide, item, 1.55, y + 0.05, 10.6, 0.50 if self_study else 0.60,
			size=19 if self_study else 13 if len(item) > 90 else 15 if len(item) > 60 else 19,
			min_size=16 if self_study else None)


def _draw_card_grid(slide, block, role, self_study):
	items = block["items"]
	numbered = block["type"] == "steps"
	columns = 3 if numbered or len(items) > 4 else 2
	if len(items) <= 2:
		columns = len(items)
	rows = (len(items) + columns - 1) // columns
	gap = 0.22
	width = (11.7 - (columns - 1) * gap) / columns
	height = 1.85 if self_study else 2.08
	top = (1.95 if rows == 2 else 2.60) if self_study else (2.05 if rows == 2 else 2.72)
	for index, item in enumerate(items):
		x = 0.8 + (index % columns) * (width + gap)
		y = top + (index // columns) * (height + 0.22)
		_panel(slide, x, y, width, height,
			fill=RGBColor(231, 240, 230) if role == "summary" else _SOFT)
		_text(slide, f"{index + 1:02d}" if numbered else "•", x + 0.22, y + 0.18,
			0.65, 0.38, size=18, color=RGBColor(41, 91, 71) if role == "summary" else _ACCENT, bold=True)
		_text(slide, item, x + 0.25, y + (0.60 if self_study else 0.65), width - 0.5,
			1.05 if self_study else 1.20,
			size=17 if len(item) > 32 else 20, min_size=16 if self_study else None, bold=role in ("hook", "summary"))


def _draw_concept(slide, block):
	items = block["items"]
	center_y = 2.05 if len(items) == 2 else 3.65
	center = (3.9, center_y, 5.53, 0.64)
	positions = [(0.8 + (index % 2) * 6.0,
		3.35 if len(items) == 2 else 1.80 if index < 2 else 4.35) for index in range(len(items))]
	# 連線先畫，保留原生可編輯圖形，且不穿過分支文字。
	for index, (x, y) in enumerate(positions):
		above = y < center_y
		connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
			Inches(center[0] + center[2] * (0.3 if index % 2 == 0 else 0.7)),
			Inches(center_y if above else center_y + center[3]),
			Inches(x + 2.85), Inches(y + 1.80 if above else y))
		connector.line.color.rgb = _ACCENT
		connector.line.width = Pt(2)
	_panel(slide, *center, fill=_ACCENT)
	_text(slide, block["center"], center[0] + 0.12, center_y + 0.02,
		center[2] - 0.24, center[3] - 0.04, size=22, min_size=16, color=_WHITE, bold=True)
	for item, (x, y) in zip(items, positions):
		_panel(slide, x, y, 5.7, 1.80)
		_text(slide, item["title"], x + 0.22, y + 0.04, 5.26, 0.59,
			size=20, min_size=16, color=_ACCENT, bold=True)
		_text(slide, item["text"], x + 0.22, y + 0.64, 5.26, 1.14,
			size=18, min_size=16)


def _draw_block(slide, block, role, self_study):
	kind = block["type"]
	if kind in ("bullets", "steps"):
		if max(len(item) for item in block["items"]) <= 48:
			_draw_card_grid(slide, block, role, self_study)
		else:
			_draw_vertical_list(slide, block, self_study)
	elif kind == "code":
		_panel(slide, 0.8, 1.95, 11.7, 4.15 if self_study else 4.5, fill=RGBColor(46, 52, 47))
		_text(slide, block["language"].upper(), 1.1, 2.15, 3.0, 0.36, size=16 if self_study else 12, min_size=16 if self_study else None, color=RGBColor(212, 173, 126), bold=True)
		lines = "\n".join(block["lines"])
		_text(slide, lines, 1.1, 2.65, 11.0, 3.35 if self_study else 3.70,
			size=18 if self_study else 12 if len(block["lines"]) > 12 else 14 if len(block["lines"]) > 9 else 17,
			min_size=16 if self_study else None,
			color=_WHITE, font="Consolas")
	elif kind == "compare":
		for x, side in ((0.8, block["left"]), (6.8, block["right"])):
			_panel(slide, x, 1.95, 5.7, 4.15 if self_study else 4.5)
			_text(slide, side["title"], x + 0.25, 2.2, 5.1, 0.5, size=22, color=_ACCENT, bold=True)
			for index, item in enumerate(side["items"]):
				_text(slide, "• " + item, x + 0.3, 2.88 + index * 0.79, 5.0, 0.70 if self_study else 0.72,
					size=18 if self_study else 12 if len(item) > 90 else 14 if len(item) > 60 else 18,
					min_size=16 if self_study else None)
	elif kind == "concept":
		_draw_concept(slide, block)


def export_pptx(value):
	"""從已驗證 Deck v1 產生 PPTX bytes，不寫入公開檔案。"""
	deck = canonical_deck(value)
	if deck['contract_version']=='1.2':
		return _export_professional(deck)
	self_study = deck["contract_version"] == "1.1"
	prs = Presentation()
	prs.slide_width = Inches(13.333)
	prs.slide_height = Inches(7.5)
	for index, item in enumerate(deck["slides"], start=1):
		slide = prs.slides.add_slide(prs.slide_layouts[5])
		slide.background.fill.solid()
		slide.background.fill.fore_color.rgb = _BG
		_text(slide, deck["title"], 0.8, 0.34, 10.8, 0.28, size=12, color=_ACCENT, bold=True)
		_text(slide, item["title"], 0.8, 0.85, 11.8, 0.83,
			size=26 if len(item["title"]) > 50 else 32, min_size=20,
			bold=True, shape=slide.shapes.title)
		_draw_block(slide, item["blocks"][0], item["role"], self_study)
		if self_study and item["layout"] == "steps":
			_text(slide, "流程方向：依序 01 → 02 → …", 0.8, 1.70, 11.7, 0.23,
				size=11, min_size=11, color=_ACCENT)
		if "takeaway" in item:
			_panel(slide, 0.8, 6.20, 11.7, 0.90, fill=_WHITE)
			_text(slide, item["takeaway"], 1.00, 6.22, 11.3, 0.86,
				size=18, min_size=16, bold=True)
		_text(slide, f"{index} / {len(deck['slides'])}", 11.8, 7.15 if self_study else 7.02,
			0.8, 0.25, size=11, color=_ACCENT)
		slide.notes_slide.notes_text_frame.text = (
			item["speaker_notes"] + "\n\n教材原文依據：\n" + item["source_quote"]
		)
	buffer = BytesIO()
	prs.save(buffer)
	return buffer.getvalue()


def _export_professional(deck):
	"""走讀每一步各留一頁原生圖形，與線上狀態完全對應。"""
	prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
	pages=[]
	for item in deck['slides']:
		block=item['blocks'][0]
		if block['type']=='walkthrough':pages.extend((item,state) for state in block['states'])
		elif block['type']=='concept':
			from copy import deepcopy
			for start in range(0,len(block['items']),2):
				page=deepcopy(item);page['blocks'][0]['items']=block['items'][start:start+2]
				pages.append((page,None))
		else:pages.append((item,None))
	for index,(item,state) in enumerate(pages,1):
		slide=prs.slides.add_slide(prs.slide_layouts[5])
		slide.background.fill.solid();slide.background.fill.fore_color.rgb=_BG
		_text(slide,deck['title'],.8,.25,11.7,.3,size=12,min_size=12,color=_ACCENT)
		_text(slide,item['title'],.8,.8,11.7,.75,size=30,min_size=24,bold=True,shape=slide.shapes.title)
		_text(slide,item['question'],.8,1.6,11.7,.38,size=20,min_size=18,color=_ACCENT)
		if state is None:
			block=item['blocks'][0]
			if block['type']=='concept':
				_panel(slide,3.9,2.2,5.53,.66,fill=_ACCENT)
				_text(slide,block['center'],4.05,2.25,5.23,.54,size=22,min_size=18,color=_WHITE,bold=True)
				for branch_index,branch in enumerate(block['items']):
					x=.8+branch_index*6.0 if len(block['items'])>1 else 3.82
					connector=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(6.665),Inches(2.86),Inches(x+2.85),Inches(3.4))
					connector.line.color.rgb=_ACCENT
					_panel(slide,x,3.4,5.7,2.5)
					_text(slide,branch['title'],x+.22,3.55,5.26,.65,size=22,min_size=18,color=_ACCENT,bold=True)
					_text(slide,branch['text'],x+.22,4.3,5.26,1.4,size=20,min_size=16)
			else:_draw_block(slide,block,item['role'],True)
		else:
			block=item['blocks'][0]
			_panel(slide,.8,2.2,7.1,3.85,fill=RGBColor(46,52,47))
			line_height=min(.42,3.4/len(block['lines']))
			for line_number,line in enumerate(block['lines'],1):
				y=2.4+(line_number-1)*line_height
				if line_number==state['line']:_panel(slide,1.0,y,6.7,line_height,fill=RGBColor(80,87,63))
				_text(slide,f'{line_number:02d}  {line}',1.08,y+.01,6.4,line_height-.02,size=18,min_size=16,color=_WHITE,font='Consolas')
			_text(slide,state['label'],8.25,2.25,4.15,.45,size=22,min_size=18,bold=True,color=_ACCENT)
			for position,variable in enumerate(state['vars']):
				y=2.9+position*.4
				_panel(slide,8.2,y,4.3,.38)
				_text(slide,variable['name'],8.4,y+.005,1.3,.36,size=18,min_size=16)
				_text(slide,variable['value'],9.8,y+.005,2.4,.36,size=18,min_size=16,bold=True)
			out_y=2.9+len(state['vars'])*.4+.1
			_text(slide,'輸出',8.25,out_y,4.0,.35,size=18,min_size=16,color=_ACCENT)
			_text(slide,state['output'],8.25,out_y+.42,4.0,.56,size=20,min_size=16,bold=True)
		_panel(slide,.8,6.2,11.7,.9,fill=_WHITE)
		_text(slide,state['explanation'] if state else item['explanation'],1.0,6.23,11.3,.84,size=20,min_size=16)
		_text(slide,f'{index} / {len(pages)}',11.8,7.15,.8,.25,size=11,min_size=11,color=_ACCENT)
		slide.notes_slide.notes_text_frame.text=(item['explanation']+'\n\n教材原文依據：\n'+
			'\n\n'.join(item['source_quotes'])+'\n\n元件選擇：'+item['selection_reason'])
	buffer=BytesIO();prs.save(buffer);return buffer.getvalue()
