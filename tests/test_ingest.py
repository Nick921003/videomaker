import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "video_engine"))
sys.path.insert(0, os.path.join(ROOT, "tests/fixtures"))

import make_pptx
import ingest
from ingest import SUPPORTED, extract_text, lesson_id_for


def make_pdf(path, page_texts):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2*i} 0 R' for i in range(len(page_texts)))}] /Count {len(page_texts)} >>".encode(),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(page_texts):
        stream_num = 5 + 2*i
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {stream_num} 0 R >>".encode())
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(pdf)
    pdf.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]: pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer << /Root 1 0 R /Size {len(objects)+1} >>\nstartxref\n{start}\n%%EOF\n".encode())
    Path(path).write_bytes(pdf)


class TestExtractText(unittest.TestCase):
	def setUp(self):
		self.dir = tempfile.mkdtemp()
		self.addCleanup(shutil.rmtree, self.dir)

	def path(self, name):
		return os.path.join(self.dir, name)

	def test_md_原文照回(self):
		p = self.path("a.md")
		with open(p, "w", encoding="utf-8") as f:
			f.write("# 標題\n\n內文一行")
		self.assertEqual(extract_text(p), "# 標題\n\n內文一行")

	def test_txt_原文照回(self):
		p = self.path("a.txt")
		with open(p, "w", encoding="utf-8") as f:
			f.write("純文字")
		self.assertEqual(extract_text(p), "純文字")

	def test_pptx_每頁都在且照播放順序(self):
		p = self.path("a.pptx")
		make_pptx.make(p, [(f"第{i}頁標題", [f"第{i}頁內容"]) for i in range(1, 6)])
		out = extract_text(p)
		for i in range(1, 6):
			self.assertIn(f"第{i}頁標題", out)
		self.assertLess(out.index("第1頁標題"), out.index("第5頁標題"))

	def test_標題不會重複出現在內文列(self):
		# slide.shapes 每次迭代都給新的 proxy 物件，用 `is` 比對會失效，
		# 標題會同時被當成標題和內文各印一次
		p = self.path("a.pptx")
		make_pptx.make(p, [("唯一標題", ["內容"])])
		self.assertEqual(extract_text(p).count("唯一標題"), 1)

	def test_備忘稿要掛在正確的頁_不是照檔名編號(self):
		# 這是本 Task 最重要的一則。只有第 2、4 頁有備忘稿時，
		# 檔案裡是 notesSlide1.xml（屬於 slide2）與 notesSlide2.xml（屬於 slide4）。
		# 照檔名編號對應會把第 2 頁的備忘稿掛到第 1 頁。
		p = self.path("a.pptx")
		make_pptx.make(p,
			[(f"第{i}頁標題", [f"第{i}頁內容"]) for i in range(1, 6)],
			{2: "這句屬於第二頁", 4: "這句屬於第四頁"})
		out = extract_text(p)
		pages = out.split("## 第 ")
		self.assertNotIn("備忘稿", pages[1])
		self.assertIn("這句屬於第二頁", pages[2])
		self.assertNotIn("備忘稿", pages[3])
		self.assertIn("這句屬於第四頁", pages[4])
		self.assertNotIn("備忘稿", pages[5])

	def test_全是圖沒有文字要丟錯(self):
		p = self.path("a.pptx")
		make_pptx.make_blank(p)
		with self.assertRaises(ValueError):
			extract_text(p)

	def test_pptx_picture_only_slide_is_marked_for_vision(self):
		from PIL import Image
		from pptx import Presentation
		from pptx.util import Inches
		p = self.path("picture.pptx")
		image = self.path("diagram.png")
		Image.new("RGB", (80, 80), "white").save(image)
		prs = Presentation()
		page = prs.slides.add_slide(prs.slide_layouts[6])
		page.shapes.add_picture(image, Inches(1), Inches(1))
		prs.save(p)
		self.assertIn("[本頁只有圖片或圖表，需要圖片辨識]", extract_text(p))

	def test_不支援的副檔名要丟錯(self):
		p = self.path("a.doc")
		with open(p, "w", encoding="utf-8") as f:
			f.write("x")
		with self.assertRaises(ValueError):
			extract_text(p)

	def test_SUPPORTED_是白名單來源(self):
		self.assertEqual(SUPPORTED, (".md", ".txt", ".pptx", ".pdf"))


	def test_pdf_text_keeps_page_boundaries_and_marks_scanned_page(self):
		p = self.path("lesson.pdf")
		make_pdf(p, ["First page", "", "Third page"])
		out = extract_text(p)
		self.assertLess(out.index("First page"), out.index("Third page"))
		self.assertIn("## 第 2 頁\n[本頁沒有可抽取的文字，需要圖片辨識]", out)
		self.assertEqual(out.count("## 第 "), 3)

	def test_pdf_rejects_fake_file_and_page_limit(self):
		p = self.path("fake.pdf")
		Path(p).write_bytes(b"not a PDF")
		with self.assertRaises(ValueError): extract_text(p)
		make_pdf(p, ["one", "two"])
		with patch.object(ingest, "MAX_PDF_PAGES", 1):
			with self.assertRaises(ValueError): extract_text(p)

	def test_empty_invalid_utf8_and_oversize_text_are_rejected(self):
		p = self.path("a.md")
		for content in (b"  \n", b"\xff"):
			with open(p, "wb") as handle:
				handle.write(content)
			with self.subTest(content=content), self.assertRaises(ValueError):
				extract_text(p)
		with open(p, "wb") as handle:
			handle.write(b"hello")
		with patch.object(ingest, "MAX_NORMALIZED_BYTES", 4):
			with self.assertRaises(ValueError):
				extract_text(p)
		self.assertEqual(ingest.MAX_NORMALIZED_BYTES, 2 * 1024 * 1024)

	def test_pptx_member_count_and_uncompressed_size_checked_before_parse(self):
		p = self.path("bomb.pptx")
		with zipfile.ZipFile(p, "w") as archive:
			archive.writestr("one.xml", "x")
			archive.writestr("two.xml", "y")
		with patch.object(ingest, "MAX_PPTX_MEMBERS", 1):
			with self.assertRaises(ValueError):
				extract_text(p)
		with patch.object(ingest, "MAX_PPTX_UNCOMPRESSED", 1):
			with self.assertRaises(ValueError):
				extract_text(p)
		self.assertEqual(ingest.MAX_PPTX_MEMBERS, 2000)
		self.assertEqual(ingest.MAX_PPTX_UNCOMPRESSED, 50 * 1024 * 1024)

	def test_pptx_slide_limit(self):
		p = self.path("slides.pptx")
		make_pptx.make(p, [("一", ["一"]), ("二", ["二"]), ("三", ["三"])])
		with patch.object(ingest, "MAX_PPTX_SLIDES", 2):
			with self.assertRaises(ValueError):
				extract_text(p)
		self.assertEqual(ingest.MAX_PPTX_SLIDES, 500)


class TestLessonId(unittest.TestCase):
	def test_英數檔名維持既有行為(self):
		self.assertEqual(lesson_id_for("/x/c_string.md"), "c_string")
		self.assertEqual(lesson_id_for("/x/C-Struct Combo.pptx"), "c_struct_combo")

	def test_中文檔名不會全部壓成底線互撞(self):
		a = lesson_id_for("/x/測試教材.pptx")
		b = lesson_id_for("/x/另一份教材.pptx")
		self.assertNotEqual(a, b)
		self.assertRegex(a, r"^[a-z0-9_]+$")


if __name__ == "__main__":
	unittest.main()
