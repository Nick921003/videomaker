"""講義：一頁一個重點、頁序固定、轉換決定、文字框不越界；放不下就指名哪一頁。"""
from io import BytesIO
import json
import unittest
import zipfile

from pptx import Presentation

from studio.deck_contracts import DeckError, canonical_deck
from studio.handout import build_handout, page_numbers
from studio.pptx_export import export_pptx
from tests.fixtures.key_points import sample_key_points
from tests.fixtures.teaching_quality import MATERIAL


def parts(payload):
	with zipfile.ZipFile(BytesIO(payload)) as archive:
		return {name: archive.read(name) for name in archive.namelist() if name != 'docProps/core.xml'}


class TestHandout(unittest.TestCase):
	def test_page_order(self):
		deck = build_handout(sample_key_points(), MATERIAL)
		self.assertEqual(deck['contract_version'], '1.3')
		self.assertEqual([slide['kind'] for slide in deck['slides']], ['cover'] + ['point'] * 4 + ['mistakes'])
		self.assertEqual(deck['slides'][1]['question'], 'for 後面括號裡為什麼用分號隔成三段？')
		self.assertEqual(len(deck['slides'][-1]['items']), 3)
		self.assertEqual(deck['slides'][-1]['items'][0]['point_title'], '迴圈三要素')
		self.assertEqual(page_numbers(sample_key_points()), {'three_parts': 2, 'choose': 3, 'boundary': 4, 'infinite': 5})

	def test_no_mistakes_page_without_misconceptions(self):
		value = sample_key_points()
		for point in value['points']:
			point['misconceptions'] = []
		deck = build_handout(value, MATERIAL)
		self.assertEqual(deck['slides'][-1]['kind'], 'point')

	def test_conversion_is_deterministic(self):
		first = json.dumps(build_handout(sample_key_points(), MATERIAL), ensure_ascii=False)
		second = json.dumps(build_handout(sample_key_points(), MATERIAL), ensure_ascii=False)
		self.assertEqual(first, second)
		deck = build_handout(sample_key_points(), MATERIAL)
		self.assertEqual(parts(export_pptx(deck)), parts(export_pptx(deck)))

	def test_contract_rejects_bad_order_and_long_text(self):
		deck = build_handout(sample_key_points(), MATERIAL)
		swapped = json.loads(json.dumps(deck)); swapped['slides'][0], swapped['slides'][1] = swapped['slides'][1], swapped['slides'][0]
		with self.assertRaises(DeckError):
			canonical_deck(swapped)
		long = json.loads(json.dumps(deck)); long['slides'][1]['explanation'] = '長' * 151
		with self.assertRaises(DeckError):
			canonical_deck(long)
		extra = json.loads(json.dumps(deck)); extra['slides'][1]['narration'] = '不應存在'
		with self.assertRaises(DeckError):
			canonical_deck(extra)

	def test_every_text_box_stays_inside_the_slide(self):
		value = sample_key_points()
		value['points'][0]['explanation'] = '長' * 150
		value['points'][0]['example']['content'] = '\n'.join(f'line_{n} = value + {n};' for n in range(12))
		value['objectives'][0]['text'] = '目' * 80
		for point in value['points']:
			point['misconceptions'] = [{'wrong': '錯' * 30, 'why': '因' * 30, 'correct': '正' * 30}]
		presentation = Presentation(BytesIO(export_pptx(build_handout(value, MATERIAL))))
		width, height = presentation.slide_width, presentation.slide_height
		for number, slide in enumerate(presentation.slides, 1):
			for shape in slide.shapes:
				self.assertGreaterEqual(shape.left, 0, number)
				self.assertGreaterEqual(shape.top, 0, number)
				self.assertLessEqual(shape.left + shape.width, width, (number, shape.name))
				self.assertLessEqual(shape.top + shape.height, height, (number, shape.name))

	def test_overflow_names_the_page(self):
		# 有例子時常見錯誤只佔左欄；兩條全滿 80 字一定放不下，必須明確報錯而不是越界。
		deck = build_handout(sample_key_points(), MATERIAL)
		deck['slides'][1]['misconceptions'] = [{'wrong': '錯' * 80, 'why': '因' * 80, 'correct': '正' * 80}] * 2
		with self.assertRaisesRegex(DeckError, '講義第 2 頁（three_parts）放不下'):
			export_pptx(deck)


if __name__ == '__main__':
	unittest.main()
