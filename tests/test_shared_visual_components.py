from copy import deepcopy
from io import BytesIO
import unittest

from pptx import Presentation
from studio.deck_contracts import canonical_deck
from studio.pptx_export import export_pptx
from studio.teaching_plan import canonical_plan, plan_deck, plan_mindmap
from tests.fixtures.teaching_quality import MATERIAL, sample_plan


def pipeline():
	return {'type': 'pipeline', 'stages': [
		{'id': 'collect', 'label': '收集', 'note': '集中原始資料。'},
		{'id': 'check', 'label': '檢查', 'note': '核對資料內容。'},
		{'id': 'publish', 'label': '發佈', 'note': '提供通過檢查的資料。'}]}


def related_plan():
	plan = sample_plan()
	plan['contract_version'] = '1.1'
	for unit in plan['units']:
		unit.update(concept_label=unit['id'], parent_id='root', relation='主題', relation_quote=unit['source_quotes'][0])
	plan['units'][2].update(parent_id='three_parts', relation='運作例子')
	plan['units'][4].update(parent_id='choose', relation='邊界檢查')
	return plan


class TestSharedVisualComponents(unittest.TestCase):
	def test_pipeline_data_matches_existing_board_component(self):
		from studio.visual_components import pipeline_object
		block = pipeline()
		obj = pipeline_object(block, 'lesson_flow')
		self.assertEqual(obj, {'id': 'lesson_flow', 'type': 'pipeline', 'stages': block['stages']})
		obj['stages'][0]['note'] = '修改'
		self.assertNotEqual(obj['stages'], block['stages'])

	def test_pipeline_rejects_duplicate_ids_unbounded_text_and_extra_code(self):
		from studio.visual_components import pipeline_object
		for kind in ('duplicate', 'count', 'text', 'extra', 'identifier'):
			block = pipeline()
			if kind == 'duplicate': block['stages'][1]['id'] = 'collect'
			elif kind == 'count': block['stages'] *= 3
			elif kind == 'text': block['stages'][0]['note'] = '字' * 81
			elif kind == 'extra': block['html'] = '<script>bad()</script>'
			else: block['stages'][0]['id'] = '../../bad'
			with self.subTest(kind=kind), self.assertRaises(ValueError): pipeline_object(block, 'flow')

	def test_new_mindmap_preserves_concept_parents_and_relation_evidence(self):
		plan = canonical_plan(related_plan(), MATERIAL)
		mindmap = plan_mindmap(plan)
		edge = next(edge for edge in mindmap['edges'] if edge['to'] == 'trace')
		self.assertEqual(edge, {'from': 'three_parts', 'to': 'trace', 'label': '運作例子',
			'source_quote': plan['units'][2]['relation_quote']})
		self.assertEqual(next(e['from'] for e in mindmap['edges'] if e['to'] == 'boundary'), 'choose')
		self.assertEqual(plan_deck(plan)['contract_version'], '1.2')
		self.assertEqual(next(n['label'] for n in mindmap['nodes'] if n['id'] == 'trace'), 'trace')
		self.assertNotEqual(plan_deck(plan)['slides'][2]['title'], 'trace')

	def test_new_relations_reject_cycles_missing_parents_and_unsupported_quotes(self):
		for kind in ('cycle', 'missing', 'self', 'quote', 'blank'):
			plan = related_plan()
			if kind == 'cycle':
				plan['units'][1]['parent_id'] = 'trace'
			elif kind == 'missing': plan['units'][2]['parent_id'] = 'missing'
			elif kind == 'self': plan['units'][2]['parent_id'] = 'trace'
			elif kind == 'quote': plan['units'][2]['relation_quote'] = '原文沒有這段'
			else: plan['units'][2]['relation'] = ''
			with self.subTest(kind=kind), self.assertRaises(ValueError): canonical_plan(plan, MATERIAL)

	def test_legacy_plan_and_deck_keep_existing_behavior(self):
		plan = canonical_plan(sample_plan(), MATERIAL)
		self.assertEqual(next(e['from'] for e in plan_mindmap(plan)['edges'] if e['to'] == 'trace'), 'root')
		self.assertEqual(len(plan_deck(plan)['slides']), len(plan['units']))

	def test_pipeline_exports_native_shapes_for_each_reading_stage(self):
		plan = sample_plan()
		plan['units'][0]['visual'] = pipeline()
		deck = canonical_deck(plan_deck(canonical_plan(plan, MATERIAL)))
		prs = Presentation(BytesIO(export_pptx(deck)))
		for index, stage in enumerate(pipeline()['stages']):
			slide = prs.slides[index]
			texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
			self.assertIn(stage['label'], texts)
			self.assertIn(stage['note'], texts)
			self.assertEqual(sum(shape.shape_type == 9 for shape in slide.shapes), 2)
			self.assertFalse(any(shape.shape_type == 13 for shape in slide.shapes))
			self.assertEqual(sum(stage['note'] == text for text in texts), 1)
			for shape in slide.shapes:
				self.assertGreaterEqual(shape.left, 0)
				self.assertLessEqual(shape.left + shape.width, prs.slide_width)
				self.assertLessEqual(shape.top + shape.height, prs.slide_height)

	def test_pipeline_dense_notes_are_rejected_without_publishing(self):
		plan = sample_plan()
		block = pipeline()
		block['stages'] += [dict(id='four', label='第四階段', note='說' * 80), dict(id='five', label='第五階段', note='說' * 80)]
		plan['units'][0]['visual'] = block
		# Either the bounded layout exports at >=16pt or explicitly rejects capacity.
		try:
			prs = Presentation(BytesIO(export_pptx(plan_deck(canonical_plan(plan, MATERIAL)))))
		except ValueError as exc:
			self.assertIn('最低可讀字級', str(exc))
		else:
			for slide in list(prs.slides)[:5]:
				for shape in slide.shapes:
					if shape.has_text_frame and shape.top > 1524000 and shape.top < 6400800:
						self.assertTrue(all(run.font.size.pt >= 16 for p in shape.text_frame.paragraphs for run in p.runs))
