from io import BytesIO
import unittest

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from studio.pptx_export import export_pptx
from tests.test_studio_deck_contracts import sample_deck, self_study_deck


class TestPptxExport(unittest.TestCase):
	def test_legacy_long_points_keep_original_readable_text_area(self):
		for length in (49, 60, 70, 120):
			with self.subTest(length=length):
				deck = sample_deck()
				text = "字" * length
				deck["slides"][0]["blocks"][0]["items"] = [text]
				prs = Presentation(BytesIO(export_pptx(deck)))
				shape = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text == text)
				self.assertEqual(shape.height, Inches(0.60))
				self.assertLessEqual(shape.top + shape.height, Inches(7.02))

	def test_legacy_dense_layouts_still_export_with_original_geometry(self):
		deck = sample_deck()
		deck["slides"][0]["blocks"][0]["items"] = ["長" * 120] * 5
		deck["slides"][1]["blocks"][0]["lines"] = ["print('" + "a" * 60 + "')"] * 16
		for side in ("left", "right"):
			deck["slides"][2]["blocks"][0][side]["items"] = ["比" * 120] * 4
		deck["slides"][3]["blocks"][0]["items"] = ["步" * 120] * 6
		prs = Presentation(BytesIO(export_pptx(deck)))
		for source, slide in zip(deck["slides"], prs.slides):
			texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
			self.assertIn(source["title"], texts)
			page = next(s for s in slide.shapes if s.has_text_frame and " / " in s.text)
			for shape in slide.shapes:
				self.assertLessEqual(shape.top + shape.height, prs.slide_height)
				self.assertLessEqual(shape.left + shape.width, prs.slide_width)
				if shape.shape_id != page.shape_id:
					self.assertLessEqual(shape.top + shape.height, page.top)
		self.assertFalse(any(s.has_text_frame and "流程方向" in s.text for s in prs.slides[3].shapes))

	def test_concept_and_takeaway_are_native_visible_and_separate_from_page_number(self):
		deck = self_study_deck()
		block = deck["slides"][0]["blocks"][0]
		block["center"] = "字" * 40
		block["items"] = [{"title": str(i) + "題" * 31, "text": str(i) + "說" * 79} for i in range(4)]
		for source in deck["slides"]:
			source["takeaway"] = "解" * 120
		prs = Presentation(BytesIO(export_pptx(deck)))
		concept = prs.slides[0]
		texts = [shape.text for shape in concept.shapes if shape.has_text_frame]
		self.assertIn(block["center"], texts)
		for item in block["items"]:
			self.assertIn(item["title"], texts)
			self.assertIn(item["text"], texts)
		self.assertEqual(sum(shape.shape_type == 9 for shape in concept.shapes), 4)
		for source, slide in zip(deck["slides"], prs.slides):
			takeaway = next(s for s in slide.shapes if s.has_text_frame and source["takeaway"] in s.text)
			page = next(s for s in slide.shapes if s.has_text_frame and " / " in s.text)
			self.assertLessEqual(takeaway.top + takeaway.height, page.top)
			for shape in slide.shapes:
				self.assertGreaterEqual(shape.left, 0)
				self.assertGreaterEqual(shape.top, 0)
				self.assertLessEqual(shape.left + shape.width, prs.slide_width)
				self.assertLessEqual(shape.top + shape.height, prs.slide_height)
				if shape.top < Inches(6.20):
					self.assertLessEqual(shape.top + shape.height, Inches(6.20))

	def test_six_long_steps_fit_or_fail_explicitly_and_never_cover_takeaway(self):
		deck = self_study_deck()
		deck["slides"][3]["blocks"][0]["items"] = ["長" * 120] * 6
		with self.assertRaisesRegex(ValueError, "最低可讀字級"):
			export_pptx(deck)

	def test_steps_show_editable_flow_direction(self):
		prs = Presentation(BytesIO(export_pptx(self_study_deck())))
		self.assertTrue(any(s.has_text_frame and "→" in s.text for s in prs.slides[3].shapes))

	def test_roundtrip_preserves_native_text_and_speaker_notes(self):
		deck = sample_deck()
		content = export_pptx(deck)
		self.assertIsInstance(content, bytes)
		prs = Presentation(BytesIO(content))
		self.assertEqual(len(prs.slides), len(deck["slides"]))
		self.assertGreater(prs.slide_width, prs.slide_height)
		for source, rendered in zip(deck["slides"], prs.slides):
			texts = [shape.text for shape in rendered.shapes if shape.has_text_frame]
			self.assertIn(source["title"], texts)
			notes = rendered.notes_slide.notes_text_frame.text
			self.assertIn(source["speaker_notes"], notes)
			self.assertIn(source["source_quote"], notes)
			self.assertIn("教材原文依據", notes)
			self.assertFalse(any(shape.shape_type == 13 for shape in rendered.shapes))
		self.assertTrue(any("for n in range(3)" in shape.text for shape in prs.slides[1].shapes if shape.has_text_frame))
		self.assertTrue(any("正確" in shape.text for shape in prs.slides[2].shapes if shape.has_text_frame))

	def test_long_slide_text_is_rejected_before_hidden_auto_shrink(self):
		deck = sample_deck()
		deck["slides"][2]["blocks"][0]["left"]["items"] = ["認識概念之間的差別" * 6] * 4
		with self.assertRaisesRegex(ValueError, "最低可讀字級"):
			export_pptx(deck)

	def test_normal_slide_text_uses_explicit_readable_size(self):
		prs = Presentation(BytesIO(export_pptx(sample_deck())))
		for slide in prs.slides:
			self.assertIsNotNone(slide.shapes.title)
			for shape in slide.shapes:
				if shape.has_text_frame and shape.text.strip():
					self.assertEqual(shape.text_frame.auto_size, MSO_AUTO_SIZE.NONE)
					self.assertLessEqual(shape.top + shape.height, prs.slide_height)
					self.assertLessEqual(shape.left + shape.width, prs.slide_width)
					self.assertTrue(all(run.font.size >= Pt(10)
						for paragraph in shape.text_frame.paragraphs for run in paragraph.runs))

	def test_short_points_and_steps_have_distinct_editable_grid_layouts(self):
		deck = sample_deck()
		deck["slides"][0]["blocks"][0]["items"] = ["問題一", "問題二", "問題三", "問題四"]
		deck["slides"][3]["blocks"][0]["items"] = [f"步驟{i}" for i in range(1, 7)]
		prs = Presentation(BytesIO(export_pptx(deck)))
		for slide, labels in ((prs.slides[0], ["問題一", "問題二", "問題三", "問題四"]),
			(prs.slides[3], [f"步驟{i}" for i in range(1, 7)])):
			boxes = [next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == label) for label in labels]
			self.assertGreater(len({box.left for box in boxes}), 1)
			self.assertGreater(len({box.top for box in boxes}), 1)
			self.assertTrue(all(box.left + box.width <= prs.slide_width and
				box.top + box.height <= prs.slide_height for box in boxes))

	def test_rejects_invalid_deck_before_export(self):
		with self.assertRaises(ValueError):
			export_pptx({"title": "bad", "slides": []})
